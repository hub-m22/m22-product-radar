"""Еженедельный управленческий отчёт: сборка, сохранение, экспорт в PDF / XLSX / CSV."""
from __future__ import annotations

import csv
import io
import sqlite3
from datetime import date, datetime, timedelta

from . import config, db
from .normalize import CATEGORY_NAMES
from .signals import SIGNAL_TYPES

SEV_ORDER = {"high": 0, "medium": 1, "low": 2}


def _cat(slug):
    return CATEGORY_NAMES.get(slug or "", slug or "—")


def _data_limits(conn: sqlite3.Connection) -> dict:
    """Ограничения текущих данных: сколько замеров, есть ли история."""
    m22_runs = db.row(conn, "SELECT COUNT(DISTINCT substr(started_at,1,10)) AS d FROM source_runs WHERE source_key IN ('m22.ru','radiosync.ru') AND status IN ('ok','partial')")["d"]
    comp_runs = db.row(conn, "SELECT COUNT(DISTINCT substr(started_at,1,10)) AS d FROM source_runs WHERE source_key='competitors' AND status IN ('ok','partial')")["d"]
    trends = db.row(conn, "SELECT COUNT(*) AS n, MIN(period_start) AS a, MAX(period_start) AS b FROM demand_observations WHERE source='google_trends'")
    ws = db.row(conn, "SELECT COUNT(*) AS n FROM demand_observations WHERE source IN ('wordstat','wordstat_import')")["n"]
    first = db.row(conn, "SELECT MIN(started_at) AS t FROM source_runs")["t"]
    notes = []
    if comp_runs <= 1:
        notes.append("Цены конкурентов сняты один раз — изменения цен (рост/снижение) станут видны после второго ежедневного сбора; сравнение «M22 vs рынок» уже корректно.")
    if m22_runs <= 1:
        notes.append("Каталог M22 снят один раз — история цен M22 начнёт накапливаться со следующего сбора.")
    if trends["n"]:
        notes.append(f"Google Trends: относительный индекс (0–100), {trends['n']} точек за {trends['a']} – {trends['b']}; абсолютные объёмы даст только Wordstat.")
    else:
        notes.append("Google Trends: данных нет (лимиты источника) — динамика спроса пока не оценивается.")
    if not ws:
        notes.append("Яндекс Wordstat не подключён (нужен OAuth-токен API или ручной импорт CSV) — абсолютных объёмов спроса и сезонности по Яндексу нет.")
    return {"m22_collection_days": m22_runs, "competitor_collection_days": comp_runs, "trends_points": trends["n"], "wordstat_points": ws,
            "first_collection": first, "notes": notes,
            "what_after_history": ["Через 7 дней: изменения цен конкурентов, появление/исчезновение товаров, первые недельные сравнения.",
                                   "Через 30 дней: тренды цен по категориям, стабильные пробелы ассортимента, месячная динамика спроса.",
                                   "Через 12 месяцев (или после импорта истории Wordstat): сезонность и аномалии на годовой базе."]}


def build_weekly(conn: sqlite3.Connection, days: int = 7) -> dict:
    end = date.today()
    start = end - timedelta(days=days)
    since = start.isoformat()
    sigs = db.rows(conn, "SELECT * FROM signals WHERE status!='rejected' AND substr(created_at,1,10)>=? ORDER BY created_at DESC", (since,))
    if not sigs:  # если за неделю ничего нового — берём все актуальные
        sigs = db.rows(conn, "SELECT * FROM signals WHERE status!='rejected' ORDER BY created_at DESC")
    sigs.sort(key=lambda s: (SEV_ORDER.get(s["severity"], 3), -(s["confidence"] or 0)))
    recs = db.rows(conn, "SELECT * FROM recommendations WHERE status IN ('new','accepted','in_progress') ORDER BY priority, created_at DESC")
    hyps = db.rows(conn, "SELECT * FROM hypotheses WHERE decision_status IN ('new','research') ORDER BY created_at DESC")

    def by_type(*types):
        return [s for s in sigs if s["type"] in types]

    def brief(s):
        return {"id": s["id"], "title": s["title"], "type": SIGNAL_TYPES.get(s["type"], s["type"]), "severity": s["severity"], "confidence": s["confidence"],
                "fact_kind": s["fact_kind"], "what": s["what_happened"], "period": s["period"], "source": s["source"], "source_url": s["source_url"],
                "why": s["why_matters"], "action": s["recommended_action"], "category": _cat(s["category_slug"]), "observed_at": s["observed_at"]}

    top5 = [brief(s) for s in sigs if s["type"] != "source_error"][:5]
    m22_changes = [brief(s) for s in by_type("m22_price_above_market", "m22_price_below_market", "cross_site_discrepancy", "competitor_price_change") if s["m22_product_id"]][:10]
    price_changes = [brief(s) for s in by_type("competitor_price_change")][:15]
    new_products = [brief(s) for s in by_type("product_appeared", "new_kit_solution", "multi_competitor_product", "new_category")][:15]
    demand = [brief(s) for s in by_type("demand_change", "demand_anomaly", "category_growth_existing", "category_growth_gap")][:15]
    gaps = [brief(s) for s in by_type("new_category", "multi_competitor_product", "category_growth_gap") if "отсутствует" in s["title"] or "нет у M22" in s["title"] or "у M22 её нет" in s["title"] or s["type"] == "new_category"][:15]
    risks = [brief(s) for s in sigs if s["severity"] == "high"][:8] + [brief(s) for s in by_type("source_error")][:3]
    owner_decisions = [{"id": r["id"], "title": r["title"], "action": r["action"], "priority": r["priority"], "confidence": r["confidence"], "basis": r["basis"], "due": r["due_date"]}
                       for r in recs if r["priority"] == "P1" or (r["owner"] or "") == "Собственник"][:8]
    next_week = [{"id": r["id"], "title": r["title"], "action": r["action"], "priority": r["priority"], "owner": r["owner"], "due": r["due_date"], "confidence": r["confidence"],
                  "expected_effect": r["expected_effect"], "sources": db.uj(r["sources_json"], [])} for r in recs][:12]
    hyp_cards = [{"id": h["id"], "title": h["title"], "description": h["description"], "next_step": h["next_step"], "missing_data": h["missing_data"], "status": h["decision_status"]} for h in hyps][:8]
    stats = {
        "m22_products": db.row(conn, "SELECT COUNT(*) AS n FROM m22_products WHERE is_active=1 AND in_scope=1")["n"],
        "competitors": db.row(conn, "SELECT COUNT(*) AS n FROM competitors WHERE is_active=1")["n"],
        "competitor_products": db.row(conn, "SELECT COUNT(*) AS n FROM competitor_products WHERE is_active=1")["n"],
        "priced": db.row(conn, "SELECT COUNT(*) AS n FROM competitor_products WHERE is_active=1 AND price IS NOT NULL")["n"],
        "signals_total": len(sigs), "signals_high": sum(1 for s in sigs if s["severity"] == "high"),
        "recommendations": len(recs), "hypotheses": len(hyps),
        "sources_ok": db.row(conn, "SELECT COUNT(*) AS n FROM sources WHERE status='ok'")["n"],
        "sources_error": db.row(conn, "SELECT COUNT(*) AS n FROM sources WHERE status='error'")["n"],
    }
    return {"kind": "weekly", "period_start": start.isoformat(), "period_end": end.isoformat(), "generated_at": db.now_iso(), "stats": stats,
            "top5": top5, "m22_changes": m22_changes, "price_changes": price_changes, "new_products": new_products, "demand": demand, "gaps": gaps,
            "hypotheses": hyp_cards, "risks": risks, "owner_decisions": owner_decisions, "next_week": next_week, "limits": _data_limits(conn)}


def save(conn: sqlite3.Connection, rep: dict) -> int:
    cur = conn.execute("INSERT INTO reports(kind, period_start, period_end, content_json) VALUES(?,?,?,?)", (rep["kind"], rep["period_start"], rep["period_end"], db.j(rep)))
    conn.commit()
    return int(cur.lastrowid)


SECTIONS = [("top5", "1. Пять главных изменений"), ("m22_changes", "2. Что произошло с товарами M22"), ("price_changes", "3. Изменения цен конкурентов"),
            ("new_products", "4. Новые товары и решения"), ("demand", "5. Изменения спроса"), ("gaps", "6. Товарные пробелы"), ("hypotheses", "7. Новые продуктовые гипотезы"),
            ("risks", "8. Главные риски"), ("owner_decisions", "9. Решения, требующие участия собственника"), ("next_week", "10. Конкретные действия на следующую неделю")]


def to_rows(rep: dict) -> list[dict]:
    rows = []
    for key, title in SECTIONS:
        for it in rep.get(key, []):
            rows.append({"section": title, "title": it.get("title"), "what": it.get("what") or it.get("description") or it.get("action"), "period": it.get("period") or it.get("due"),
                         "source": it.get("source") or "; ".join(s.get("name", "") for s in (it.get("sources") or [])), "source_url": it.get("source_url"),
                         "confidence": it.get("confidence"), "severity": it.get("severity") or it.get("priority"), "why": it.get("why") or it.get("basis"),
                         "action": it.get("action") or it.get("next_step"), "owner": it.get("owner")})
    return rows


def to_csv(rep: dict) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=["section", "title", "what", "period", "source", "source_url", "confidence", "severity", "why", "action", "owner"], delimiter=";")
    w.writeheader()
    for r in to_rows(rep):
        w.writerow(r)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_xlsx(rep: dict) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.append([f"M22 Product Radar — еженедельный отчёт {rep['period_start']} — {rep['period_end']} (сформирован {rep['generated_at'][:16].replace('T', ' ')} UTC)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Раздел", "Заголовок", "Что произошло", "Период", "Источник", "Ссылка", "Уверенность", "Важность/приоритет", "Почему важно", "Действие", "Ответственный"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in to_rows(rep):
        ws.append([r["section"], r["title"], r["what"], r["period"], r["source"], r["source_url"], r["confidence"], r["severity"], r["why"], r["action"], r["owner"]])
    for col, width in zip("ABCDEFGHIJK", (28, 50, 60, 18, 24, 30, 11, 12, 50, 60, 18)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2 = wb.create_sheet("Ограничения данных")
    for n in rep["limits"]["notes"]:
        ws2.append([n])
    ws2.append([])
    for n in rep["limits"]["what_after_history"]:
        ws2.append([n])
    ws2.column_dimensions["A"].width = 120
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _register_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in (config.PDF_FONT_PATH, "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/segoeui.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/dejavu/DejaVuSans.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"):
        try:
            pdfmetrics.registerFont(TTFont("RadarFont", path))
            return "RadarFont"
        except Exception:  # noqa: BLE001
            continue
    return "Helvetica"


def to_pdf(rep: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _register_font()
    h1 = ParagraphStyle("h1", fontName=font, fontSize=15, leading=19, spaceAfter=6)
    h2 = ParagraphStyle("h2", fontName=font, fontSize=12, leading=15, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#1f3a5f"))
    body = ParagraphStyle("b", fontName=font, fontSize=9, leading=12)
    small = ParagraphStyle("s", fontName=font, fontSize=8, leading=10, textColor=colors.HexColor("#555555"))
    esc = lambda t: (str(t or "")).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    story = [Paragraph(f"M22 Product Radar — еженедельный отчёт {rep['period_start']} — {rep['period_end']}", h1),
             Paragraph(f"Сформирован {esc(rep['generated_at'][:16].replace('T', ' '))} UTC. Товаров M22 в контуре: {rep['stats']['m22_products']}; конкурентов: {rep['stats']['competitors']}; "
                       f"предложений конкурентов с ценой: {rep['stats']['priced']}; сигналов: {rep['stats']['signals_total']} (высокой важности: {rep['stats']['signals_high']}).", small)]
    for key, title in SECTIONS:
        items = rep.get(key, [])
        story.append(Paragraph(title, h2))
        if not items:
            story.append(Paragraph("Нет данных за период.", small))
            continue
        for it in items:
            head = esc(it.get("title"))
            meta = []
            if it.get("severity") or it.get("priority"):
                meta.append(f"важность: {esc(it.get('severity') or it.get('priority'))}")
            if it.get("confidence") is not None:
                meta.append(f"уверенность: {float(it['confidence']):.0%}")
            if it.get("period") or it.get("due"):
                meta.append(f"период/срок: {esc(it.get('period') or it.get('due'))}")
            if it.get("source"):
                meta.append(f"источник: {esc(it.get('source'))}")
            story.append(Paragraph(f"<b>{head}</b>", body))
            what = it.get("what") or it.get("description") or ""
            if what:
                story.append(Paragraph(esc(what), body))
            if it.get("why") or it.get("basis"):
                story.append(Paragraph("Почему важно: " + esc(it.get("why") or it.get("basis")), body))
            if it.get("action") or it.get("next_step"):
                story.append(Paragraph("Действие: " + esc(it.get("action") or it.get("next_step")), body))
            if meta:
                story.append(Paragraph(" · ".join(meta), small))
            story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("Ограничения текущих данных", h2))
    for n in rep["limits"]["notes"]:
        story.append(Paragraph("• " + esc(n), body))
    story.append(Paragraph("Что появится после накопления истории", h2))
    for n in rep["limits"]["what_after_history"]:
        story.append(Paragraph("• " + esc(n), body))
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title="M22 Product Radar — отчёт")
    doc.build(story)
    return out.getvalue()
